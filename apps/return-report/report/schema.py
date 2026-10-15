"""录入模板的表结构，字段名与《数据字段清单_20260928》一致。

kind: 字段 = 字段清单原有字段；关联 = 字段清单中由页面上下文隐含、录入表里需要显式填写的关联键；
      新增 = 生成报告需要、字段清单中尚未定义的建议新增字段。
"""
from __future__ import annotations

from dataclasses import dataclass

import openpyxl
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


@dataclass
class Col:
    name: str
    required: bool = False
    type: str = "文本"
    desc: str = ""
    kind: str = "字段"
    width: int = 14


def F(name, req="否", typ="文本", desc="", kind="字段", width=14):
    return Col(name, req == "是", typ, desc, kind, width)


THIRD_PARTY = [
    F("ISIN code", "是", desc="基金份额（class share）的唯一识别码"),
    F("LEI code", "是", desc="基金的识别码，不同份额共用同一个LEI code"),
    F("基金简称", "是", desc="系统内的唯一名称，申赎记录、净值、收益率都按它关联", width=16),
    F("基金全称", "是", width=24),
    F("成立日期", "是", "日期", "YYYY/MM/DD", width=12),
    F("到期日", typ="日期", desc="有固定期限的产品填写；不早于成立日期", width=12),
    F("基金类型", "是", "下拉", "开放式/封闭式"),
    F("币种", "是", "下拉", "JPY / USD / Euro", width=8),
    F("股份类别", "是", desc="报告 Table-3 显示为「基金简称 - 股份类别」"),
    F("运行状态", "是", "下拉", "已上架/已上架未投/待上架"),
    F("基金规模", "是", "数值", "单位：亿美元"),
    F("申购开放日", "是", "单选", "频率：日、周、月、季；日期：第一个工作日、第一个自然日、最后一个工作日、最后一个自然日"),
    F("赎回开放日", "是", "单选", "同申购开放日"),
    F("申购文件截止日期", typ="数值+单选", desc="申购开放日前【x】天/工作日"),
    F("申购款截止日期", typ="数值+单选", desc="申购开放日前【x】天/工作日"),
    F("赎回通知期", "是", "数值+单选", "【x】天/工作日"),
    F("赎回款到账时间", typ="数值+单选", desc="【x】天/工作日"),
    F("最低初次申购金额", "是", "数值+单选", "单位：【】对应币种"),
    F("最低追加申购金额", "是", "数值+单选", "单位：【】对应币种"),
    F("最低赎回金额", "是", "数值+单选", "单位：【】对应币种"),
    F("赎回金额限制", "是", "数值+单选", "Gate：【x】天/工作日+ILG/FLG"),
    F("追踪负责人", "是", "单选", "产品组所有人"),
    F("约会、业绩负责人", "是"),
    F("Amrita Feeder", "是", "下拉", "Feeder名/否"),
    F("基金策略", "是", "下拉", "多策略、量化、固收、股票多空、股票多头、商品、宏观、波动率、CTA、事件驱动、加密、私募风投"),
    F("基金策略详情", "是", desc="报告 Table-1 显示为「基金简称（基金策略详情）」", width=18),
    F("配置策略", "是", "下拉", "均衡/进攻/机会/现金管理；报告中显示为「均衡型」等"),
    F("投资地区", "是", "下拉", "全球/美国/欧洲/大中华/泛亚洲/韩国/日本/新兴市场"),
    F("申购费率", typ="多行文本"),
    F("赎回费率", "是", "多行文本", "如有软锁定才有提前赎回费"),
    F("管理费率", "是", "多行文本"),
    F("业绩费率", "是", "多行文本"),
    F("管理人", "是"),
    F("基金经理", "是"),
    F("投资团队邮箱"),
    F("投资团队电话"),
    F("投资者关系", "是"),
    F("IR邮箱1", "是"),
    F("IR邮箱2"),
    F("IR电话"),
    F("行政管理人"),
    F("行政管理人邮箱1"),
    F("行政管理人邮箱2"),
    F("行政管理人电话"),
]

OWN_FUND = [
    F("ISIN code", "是"), F("LEI code", "是"), F("基金简称", "是", width=16), F("基金全称", "是", width=24),
    F("成立日期", "是", "日期", width=12), F("到期日", typ="日期", width=12),
    F("产品级别", "是", "下拉", "一级/二级"), F("币种", "是", "下拉", "JPY / USD", width=8),
    F("运行状态", "是", "下拉", "已上架/待补充"), F("基金规模", "是", "数值", "单位：亿美元"),
    F("申购开放日", "是", "单选"), F("赎回开放日", "是", "单选"),
    F("申购文件截止日期", typ="文本", desc="YYYY/MM/DD"), F("申购款截止日期", typ="文本", desc="YYYY/MM/DD"),
    F("赎回通知期", "是"), F("赎回款到账时间"),
    F("最低初次申购金额", "是", "数值", "单位：百万美元"), F("最低追加申购金额", "是", "数值", "单位：百万美元"),
    F("最低赎回金额", "是", "数值", "单位：百万美元"), F("赎回金额限制", "是"),
    F("基金策略", "是", "下拉"), F("基金策略详情", "是", "下拉"),
    F("配置策略", "是", "下拉", "均衡/进攻/机会"), F("投资地区", "是", "下拉"),
    F("申购费率", typ="多行文本"), F("赎回费率", "是", "多行文本"), F("管理费率", "是", "多行文本"),
    F("银行名称"), F("账号"), F("账号种类"), F("现金余额"),
]

OWN_FUND_COMPONENTS = [
    F("自营基金简称", "是", desc="对应「自营基金」表的基金简称", kind="关联", width=16),
    F("成分子基金", "是", "下拉", "从未删除的三方基金中选择", width=16),
    F("份额", typ="数值"), F("权重", typ="数值（%）", desc="自动计算"), F("日期", typ="日期", width=12), F("备注"),
]

CLIENT = [
    F("客户编号", "是", desc="AM + 4 位数字，如 AM0601"),
    F("投资主体名称", "是", desc="个人或机构的法定名称；报告标题使用", width=18),
    F("客户类型", "是", "下拉", "Key Account / General / Single Fund Account / FF / Internal / Strategy / LPF Client / Other Accounts-no fee / Dormant Account / Legacy"),
    F("客户类型2", "是", "下拉", "Institutional / Corporate / Individual"),
    F("常用名"),
    F("账户状态", "是", "下拉", "Active / Fully redeemed"),
    F("签署协议", typ="多选", desc="IAA / LOA(EAM) / LOA(AEA)，可多选；未签署填 n/a"),
    F("产品", typ="多选", desc="OFC / 3PFund / EAM / AEA，可多选"),
    F("首次投资 OFC 日期", typ="日期", desc="仅 OFC 客户填写", width=12),
    F("首次投资日期", typ="日期", width=12),
    F("定期回顾表发送情况", typ="下拉"),
    F("定期回顾表签署日期", typ="日期", width=12),
    F("最近 CDD 批准日期", typ="日期", width=12),
    F("客户风险评级", typ="下拉", desc="Low / Medium / High"),
    F("核证副本提供方"),
    F("专业投资者资产证明", typ="下拉"),
    F("PI Letter 签署日期", typ="日期", width=12),
    F("税务居民地"),
    F("Amrita 对接人", typ="下拉"),
    F("客户联系人"),
    F("定期回顾表接收邮箱"),
    F("所在城市"),
    F("通讯地址", typ="多行文本"),
    F("是否出月报", typ="下拉", desc="是/否；批量生成月报时只处理「是」的客户", kind="新增"),
]

TRANSACTIONS = [
    F("客户编号", "是", desc="对应「客户」表；字段清单中由客户页面上下文确定", kind="关联"),
    F("交易类型", "是", "单选", "初次申购 /追加申购/ 赎回"),
    F("基金", "是", "下拉", "填基金简称；赎回只能针对当前有持仓的基金", width=16),
    F("确认日期", "是", "日期", "报告按确认日期所在月份记账（视为月初发生）", width=12),
    F("金额", "是", "数值", "申购：申购金额；赎回：赎回到账金额"),
    F("币种", "是", "下拉", "CNY / USD / HKD", width=8),
    F("确认份额", "是", "数值", "大于 0；持有份额 = Σ申购份额 − Σ赎回份额"),
    F("确认净值", typ="数值", desc="4 位小数"),
    F("费率", typ="数值（%）"),
    F("费用", typ="数值"),
    F("备注", typ="多行文本", width=40),
]

BENCHMARK = [
    F("基准名称", "是", desc="报告中按本表顺序展示", width=24),
    F("基准分类", desc="手工录入"),
    F("备注", typ="多行文本", width=40),
]

BENCHMARK_NAV = [
    F("基准名称", "是", desc="对应「基准」表", kind="关联", width=24),
    F("日期", "是", "日期", "月末日期；月度收益 = 本月末 / 上月末 − 1", width=12),
    F("净值", "是", "数值", "Index Value"),
]

FUND_NAV = [
    F("基金简称", "是", desc="对应「三方基金」或「自营基金」表", kind="关联", width=16),
    F("日期", "是", "日期", "支持 2026-01-05 / 2026/1/5 / Excel 日期", width=12),
    F("单位净值", "是", "数值", "持有市值 = 持有份额 × 单位净值；没有净值的月份由 MTD 收益率推算"),
]

FUND_RETURN = [
    F("基金简称", "是", desc="对应「三方基金」或「自营基金」表", kind="关联", width=16),
    F("日期", "是", "日期", "月末日期", width=12),
    F("MTD收益率", "是", "数值", "小数，3.34% 填 0.0334；只有区间时留空并填区间上下限"),
    F("是否预估值", typ="文本", desc="是/否"),
    F("区间下限", typ="数值", desc="基金只给出区间估算时填写", kind="新增"),
    F("区间上限", typ="数值", desc="基金只给出区间估算时填写", kind="新增"),
    F("估算点值", typ="数值", desc="用于推算净值；不填则用区间中点", kind="新增"),
    F("备注", typ="文本", kind="新增", width=30),
]

SHEETS: dict[str, list[Col]] = {
    "三方基金": THIRD_PARTY,
    "自营基金": OWN_FUND,
    "自营基金_成分基金": OWN_FUND_COMPONENTS,
    "客户": CLIENT,
    "申赎记录": TRANSACTIONS,
    "基准": BENCHMARK,
    "基准净值": BENCHMARK_NAV,
    "基金净值": FUND_NAV,
    "基金收益率": FUND_RETURN,
}

PERCENT_COLS = {("基金收益率", c) for c in ("MTD收益率", "区间下限", "区间上限", "估算点值")}

@dataclass
class Tagged:
    value: object
    tag: str


def mock(v):
    return Tagged(v, "模拟")


def derived(v):
    return Tagged(v, "推算")


CELL_FILLS = {
    "推算": PatternFill("solid", fgColor="FFEB9C"),
    "模拟": PatternFill("solid", fgColor="FFC7CE"),
}
_TAG_BY_RGB = {f.fgColor.rgb[-6:]: tag for tag, f in CELL_FILLS.items()}


def tag_of(cell) -> str | None:
    rgb = cell.fill.fgColor.rgb if cell.fill and cell.fill.fill_type else None
    return _TAG_BY_RGB.get(rgb[-6:]) if isinstance(rgb, str) else None

FILLS = {
    "必填": PatternFill("solid", fgColor="1F4E79"),
    "选填": PatternFill("solid", fgColor="5B9BD5"),
    "关联": PatternFill("solid", fgColor="548235"),
    "新增": PatternFill("solid", fgColor="C55A11"),
}

LEGEND = [
    ("本模板按《数据字段清单_20260928》生成，每张表对应字段清单中的一个页面；表头字段名与字段清单一致。", None),
    ("表头颜色：", None),
    ("　深蓝 = 字段清单中的必填字段", "必填"),
    ("　浅蓝 = 字段清单中的选填字段", "选填"),
    ("　绿色 = 关联键：字段清单中由页面上下文隐含（如在客户页面里录入申赎），录入表里需要显式填写", "关联"),
    ("　橙色 = 建议新增：生成月报需要、字段清单中尚未定义", "新增"),
    ("数据单元格颜色：", None),
    ("　无底色 = 真实数据（来自对账单、基金月报、tracker 等原始资料）", None),
    ("　黄色 = 推算数据：由真实数据换算得到，如按持仓市值反推的份额和单位净值、按月度收益反推的指数点位", "cell:推算"),
    ("　红色 = 模拟数据：目前没有来源、为跑通流程临时编造的值，正式使用前必须替换", "cell:模拟"),
    ("鼠标停在表头上可以看到字段类型和取值说明。", None),
    ("字段清单中「系统自动维护」「持仓（由申赎记录实时汇总）」「文档上传」等不需要录入的字段未放入模板。", None),
]


def write_workbook(path: str, rows: dict[str, list[dict]] | None = None, notes: list[str] | None = None) -> None:
    rows = rows or {}
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "说明"
    ws.column_dimensions["A"].width = 120
    for text, key in LEGEND + [(n, None) for n in (notes or [])]:
        ws.append([text])
        if key and key.startswith("cell:"):
            ws.cell(ws.max_row, 1).fill = CELL_FILLS[key[5:]]
        elif key:
            ws.cell(ws.max_row, 1).font = Font(color="FFFFFF", bold=True)
            ws.cell(ws.max_row, 1).fill = FILLS[key]

    for sheet, cols in SHEETS.items():
        ws = wb.create_sheet(sheet)
        ws.append([c.name for c in cols])
        ws.freeze_panes = "A2"
        for i, c in enumerate(cols, 1):
            cell = ws.cell(1, i)
            key = c.kind if c.kind in ("关联", "新增") else ("必填" if c.required else "选填")
            cell.fill = FILLS[key]
            cell.font = Font(color="FFFFFF", bold=True)
            cell.alignment = Alignment(wrap_text=True, vertical="center")
            info = f"{'必填' if c.required else '选填'}｜{c.type}" + (f"\n{c.desc}" if c.desc else "")
            cell.comment = Comment(info, "字段清单", width=320, height=120)
            ws.column_dimensions[get_column_letter(i)].width = c.width
        for r in rows.get(sheet, []):
            vals = [r.get(c.name) for c in cols]
            ws.append([v.value if isinstance(v, Tagged) else v for v in vals])
            for i, v in enumerate(vals, 1):
                if isinstance(v, Tagged):
                    ws.cell(ws.max_row, i).fill = CELL_FILLS[v.tag]
        for i, c in enumerate(cols, 1):
            fmt = ("0.00%" if (sheet, c.name) in PERCENT_COLS else "yyyy-mm-dd" if c.type == "日期"
                   else "#,##0.00" if c.name in ("金额", "费用") else None)
            if fmt:
                for row in ws.iter_rows(min_row=2, min_col=i, max_col=i):
                    row[0].number_format = fmt
    wb.save(path)
