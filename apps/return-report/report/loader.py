from __future__ import annotations

import re
from datetime import date, datetime

import openpyxl

from .model import Dataset, Fund, Month, Nav, Override, Return, to_month


def normalize_name(name: str) -> str:
    """'Alpha（平台型多策略）' / 'Alpha - Class A' -> 'Alpha'"""
    s = re.split(r"[（(]", str(name))[0]
    return s.split(" - ")[0].strip()


def parse_return(v) -> Return | None:
    if v is None or v == "":
        return None
    if isinstance(v, (int, float)):
        return Return(value=float(v))
    num = r"(-?\d+(?:\.\d+)?)\s*%?"
    s = str(v).strip()
    if m := re.fullmatch(rf"{num}\s*[-~～至到]\s*{num}", s):
        lo, hi = sorted(float(x) / 100 for x in m.groups())
        return Return(lo=lo, hi=hi)
    if m := re.fullmatch(num, s):
        return Return(value=float(m.group(1)) / 100)
    raise ValueError(f"无法识别的收益率：{v!r}")


def _rows(ws):
    return [list(r) for r in ws.iter_rows(values_only=True)]


def _find(rows, prefix: str) -> int:
    for i, r in enumerate(rows):
        if len(r) > 1 and isinstance(r[1], str) and r[1].startswith(prefix):
            return i
    raise ValueError(f"tracker 中找不到 {prefix}")


def _table(rows, title: str):
    h = _find(rows, title)
    cols = {ci: to_month(v) for ci, v in enumerate(rows[h]) if isinstance(v, (datetime, date))}
    out = {}
    r = h + 1
    while r < len(rows) and rows[r][1] not in (None, "") and not str(rows[r][1]).startswith("总计"):
        out[rows[r][1]] = {m: rows[r][ci] for ci, m in cols.items()}
        r += 1
    return out, sorted(set(cols.values()))


def load_tracker(path: str) -> dict:
    """读取 Return tracker 的 T 页：Table-1 收益率、Table-2 投资金额、Table-3 净值。"""
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    rows = _rows(wb["T"])
    wb.close()

    t1, months = _table(rows, "Table-1")
    t2, _ = _table(rows, "Table-2")
    t3, _ = _table(rows, "Table-3")
    last = months[-1]

    returns, navs, flows = {}, {}, {}
    for name, vals in t1.items():
        k = normalize_name(name)
        returns[k] = {}
        for m, v in vals.items():
            r = parse_return(v)
            if r:
                r.source = "tracker"
                r.estimate = m == last and k in {normalize_name(n) for n in t3}
                returns[k][m] = r
    for name, vals in t2.items():
        flows[normalize_name(name)] = {m: float(v) for m, v in vals.items() if v}
    for name, vals in t3.items():
        k = normalize_name(name)
        navs[k] = {
            m: Nav(float(v), "tracker估算" if m == last else "tracker", estimate=m == last)
            for m, v in vals.items() if isinstance(v, (int, float))
        }
    return {"returns": returns, "navs": navs, "flows": flows, "last_month": last,
            "names": [normalize_name(n) for n in t1]}


def _sheet_dicts(wb, name):
    rows = _rows(wb[name])
    header = [str(h).strip() if h is not None else "" for h in rows[0]]
    return [dict(zip(header, r)) for r in rows[1:] if any(v not in (None, "") for v in r)]


def _month(v) -> Month:
    if isinstance(v, (datetime, date)):
        return to_month(v)
    y, m = re.findall(r"\d+", str(v))[:2]
    return (int(y), int(m))


def load_increment(path: str) -> dict:
    wb = openpyxl.load_workbook(path, data_only=True)
    info = {r["项目"]: r["值"] for r in _sheet_dicts(wb, "报告参数")}
    funds = {}
    for i, r in enumerate(_sheet_dicts(wb, "基金信息")):
        k = str(r["基金简称"]).strip()
        funds[k] = Fund(k, r["表1显示名"] or k, r["表3显示名"] or k,
                        (r["风险分类"] or "").strip(), (r["投资地区"] or "").strip(),
                        (r["基金策略"] or "").strip(), i)

    returns = []
    for r in _sheet_dicts(wb, "收益率"):
        ret = parse_return(r["MTD收益率"]) if r["MTD收益率"] not in (None, "") else None
        if ret is None and r["区间下限"] is not None:
            ret = Return(lo=float(r["区间下限"]), hi=float(r["区间上限"]))
        if ret is None:
            continue
        if r.get("估算点值") not in (None, ""):
            ret.point = float(r["估算点值"])
        ret.estimate = str(r["是否预估值"]).strip() == "是"
        ret.source = "增量"
        ret.note = r.get("备注") or ""
        returns.append((str(r["基金/基准"]).strip(), _month(r["月份"]), ret))

    navs = [(str(r["基金简称"]).strip(), _month(r["月份"]), Nav(float(r["持仓净值"]), "对账单"))
            for r in _sheet_dicts(wb, "净值")]
    flows = [(str(r["基金简称"]).strip(), _month(r["日期"]), float(r["金额"]), r.get("备注") or "")
             for r in _sheet_dicts(wb, "申赎")]
    as_of = _month(info["报告截止日"])
    return {"client": info.get("客户", "") or "", "as_of": as_of, "funds": funds,
            "returns": returns, "navs": navs, "flows": flows}


def build_dataset(tracker_path: str, increment_path: str) -> Dataset:
    t = load_tracker(tracker_path)
    inc = load_increment(increment_path)
    funds = inc["funds"]
    benchmarks = [n for n in t["names"] if n not in funds]
    ds = Dataset(inc["client"], inc["as_of"], funds, benchmarks,
                 tracker_estimate_months=[t["last_month"]])

    unknown = (set(t["returns"]) | set(t["navs"]) | set(t["flows"])) - set(funds) - set(benchmarks)
    if unknown:
        raise ValueError(f"tracker 中的基金未在「基金信息」中登记：{sorted(unknown)}")

    for k in list(funds) + benchmarks:
        ds.returns[k] = dict(t["returns"].get(k, {}))
    for k in funds:
        ds.navs[k] = dict(t["navs"].get(k, {}))
        ds.flows[k] = dict(t["flows"].get(k, {}))

    for k, m, r in inc["returns"]:
        if k not in ds.returns:
            if k in funds:
                ds.returns[k] = {}
            else:
                ds.benchmarks.append(k)
                ds.returns[k] = {}
        old = ds.returns[k].get(m)
        if old and not old.same_as(r):
            ds.overrides.append(Override("收益率", k, m, old, r, r.note))
        ds.returns[k][m] = r
    for k, m, n in inc["navs"]:
        old = ds.navs.setdefault(k, {}).get(m)
        if old and abs(old.value - n.value) > 0.5:
            ds.overrides.append(Override("净值", k, m, old, n))
        ds.navs[k][m] = n
    for k, m, amt, note in inc["flows"]:
        old = ds.flows.setdefault(k, {}).get(m)
        if old is not None and abs(old - amt) > 0.005:
            ds.overrides.append(Override("申赎", k, m, old, amt, note))
        ds.flows[k][m] = amt
    return ds
