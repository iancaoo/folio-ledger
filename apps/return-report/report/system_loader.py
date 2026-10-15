"""读取按《数据字段清单》格式填写的录入表，组装成报告计算用的 Dataset。"""
from __future__ import annotations

from collections import defaultdict
from datetime import date, datetime

import openpyxl

from .loader import parse_return
from .model import Dataset, Fund, Month, Nav, Return, to_month
from .schema import SHEETS, tag_of


class Sheet:
    def __init__(self, ws, name: str):
        self.name = name
        rows = list(ws.iter_rows())
        self.header = [str(c.value).strip() if c.value is not None else "" for c in rows[0]]
        self.rows, self.tags = [], []
        for r in rows[1:]:
            if all(c.value in (None, "") for c in r):
                continue
            self.rows.append({h: c.value for h, c in zip(self.header, r)})
            self.tags.append({h: tag_of(c) for h, c in zip(self.header, r)})

    def __iter__(self):
        return iter(zip(self.rows, self.tags))


class SystemData:
    def __init__(self, path: str):
        wb = openpyxl.load_workbook(path, data_only=True)
        self.sheets = {n: Sheet(wb[n], n) for n in SHEETS if n in wb.sheetnames}

    def clients_for_report(self) -> list[str]:
        return [str(r["客户编号"]).strip() for r, _ in self.sheets["客户"]
                if str(r.get("是否出月报") or "").strip() == "是"]

    def latest_month(self) -> Month:
        return max(_month(r["日期"]) for r, _ in self.sheets["基金收益率"])


def _month(v) -> Month:
    if isinstance(v, (datetime, date)):
        return to_month(v)
    s = str(v).replace("/", "-")
    y, m = s.split("-")[:2]
    return (int(y), int(m))


def _num(v) -> float | None:
    return None if v in (None, "") else float(v)


def load_client(sd: SystemData, client_code: str, as_of: Month | None = None) -> Dataset:
    as_of = as_of or sd.latest_month()
    prov: dict[tuple[str, str, str], int] = defaultdict(int)

    def track(sheet: str, tags: dict):
        for col, tag in tags.items():
            if tag:
                prov[(sheet, col, tag)] += 1

    client = next(((r, t) for r, t in sd.sheets["客户"] if str(r["客户编号"]).strip() == client_code), None)
    if client is None:
        raise ValueError(f"客户表中没有客户编号 {client_code}")
    track("客户", {k: v for k, v in client[1].items() if k in ("客户编号", "投资主体名称")})
    name = client[0].get("常用名") or client[0]["投资主体名称"]

    txns = [(r, t) for r, t in sd.sheets["申赎记录"]
            if str(r["客户编号"]).strip() == client_code and _month(r["确认日期"]) <= as_of]
    held = {str(r["基金"]).strip() for r, _ in txns}

    funds, missing_required = {}, []
    required = [c.name for c in SHEETS["三方基金"] if c.required]
    order = 0
    for sheet in ("三方基金", "自营基金"):
        for r, t in sd.sheets.get(sheet, []):
            k = str(r["基金简称"]).strip()
            if k not in held:
                continue
            track(sheet, {c: t[c] for c in ("基金简称", "基金策略详情", "股份类别", "配置策略", "投资地区", "基金策略")})
            detail, share = r.get("基金策略详情"), r.get("股份类别")
            role = str(r.get("配置策略") or "").strip()
            funds[k] = Fund(k, f"{k}（{detail}）" if detail else k, f"{k} - {share}" if share else k,
                            role if not role or role.endswith("型") else role + "型",
                            str(r.get("投资地区") or "").strip(), str(r.get("基金策略") or "").strip(), order)
            order += 1
            if sheet == "三方基金":
                missing_required += [(sheet, k, c) for c in required if r.get(c) in (None, "")]
    unknown = held - set(funds)
    if unknown:
        raise ValueError(f"申赎记录中的基金未在基金表中登记：{sorted(unknown)}")

    flows: dict[str, dict[Month, float]] = {k: defaultdict(float) for k in funds}
    unit_moves: dict[str, list[tuple[Month, float]]] = {k: [] for k in funds}
    for r, t in txns:
        k, m = str(r["基金"]).strip(), _month(r["确认日期"])
        sign = -1 if str(r["交易类型"]).strip() == "赎回" else 1
        flows[k][m] += sign * float(r["金额"])
        unit_moves[k].append((m, sign * float(r["确认份额"])))
        track("申赎记录", t)

    unit_nav: dict[str, dict[Month, tuple[float, dict]]] = {k: {} for k in funds}
    for r, t in sd.sheets["基金净值"]:
        k = str(r["基金简称"]).strip()
        if k in funds and _month(r["日期"]) <= as_of:
            unit_nav[k][_month(r["日期"])] = (float(r["单位净值"]), t)

    navs: dict[str, dict[Month, Nav]] = {k: {} for k in funds}
    for k in funds:
        for m, (px, t) in unit_nav[k].items():
            units = sum(q for mm, q in unit_moves[k] if mm <= m)
            if units > 1e-9:
                navs[k][m] = Nav(units * px, "对账单")
                track("基金净值", t)

    returns: dict[str, dict[Month, Return]] = {k: {} for k in funds}
    for r, t in sd.sheets["基金收益率"]:
        k, m = str(r["基金简称"]).strip(), _month(r["日期"])
        if k not in funds or m > as_of:
            continue
        ret = parse_return(r["MTD收益率"]) if r.get("MTD收益率") not in (None, "") else None
        if ret is None and r.get("区间下限") not in (None, ""):
            ret = Return(lo=_num(r["区间下限"]), hi=_num(r["区间上限"]))
        if ret is None:
            continue
        ret.point = _num(r.get("估算点值"))
        ret.estimate = str(r.get("是否预估值") or "").strip() == "是"
        ret.note = r.get("备注") or ""
        ret.source = "基金收益率"
        returns[k][m] = ret
        track("基金收益率", t)

    benchmarks = [str(r["基准名称"]).strip() for r, _ in sd.sheets["基准"]]
    levels: dict[str, list[tuple[Month, float]]] = defaultdict(list)
    for r, t in sd.sheets["基准净值"]:
        b, m = str(r["基准名称"]).strip(), _month(r["日期"])
        if b in benchmarks and m <= as_of:
            levels[b].append((m, float(r["净值"])))
            track("基准净值", t)
    for b in benchmarks:
        pts = sorted(levels[b])
        returns[b] = {m: Return(value=v / pv - 1, source="基准净值") for (pm, pv), (m, v) in zip(pts, pts[1:])}

    ds = Dataset(str(name), as_of, funds, benchmarks, returns, navs,
                 {k: dict(v) for k, v in flows.items()}, client_code=client_code)
    ds.provenance = dict(prov)
    ds.missing_required = missing_required
    return ds
